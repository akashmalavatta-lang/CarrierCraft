"""
CC — Career Craft
Resume Parser: Extract raw text from PDF and DOCX files.
Uses PyMuPDF for PDF and python-docx for DOCX.
Gracefully handles corrupt or unsupported files.
"""
import os
import re
import logging
from typing import Optional

logger = logging.getLogger(__name__)


def extract_text_from_pdf(file_path: str) -> Optional[str]:
    """Extract text from a PDF using PyMuPDF (fitz)."""
    try:
        import fitz  # PyMuPDF
        doc = fitz.open(file_path)
        pages = []
        for page in doc:
            pages.append(page.get_text("text"))
        doc.close()
        return "\n".join(pages)
    except ImportError:
        logger.warning("PyMuPDF not installed — trying pdfplumber fallback")
        return _pdf_fallback(file_path)
    except Exception as e:
        logger.error(f"PDF extraction failed for {file_path}: {e}")
        return None


def _pdf_fallback(file_path: str) -> Optional[str]:
    """Fallback PDF extraction using pdfplumber if PyMuPDF unavailable."""
    try:
        import pdfplumber
        with pdfplumber.open(file_path) as pdf:
            return "\n".join(
                page.extract_text() or "" for page in pdf.pages
            )
    except Exception as e:
        logger.error(f"PDF fallback failed: {e}")
        return None


def extract_text_from_docx(file_path: str) -> Optional[str]:
    """Extract text from a DOCX file using python-docx."""
    try:
        from docx import Document
        doc = Document(file_path)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        # Also extract table text
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        paragraphs.append(cell.text)
        return "\n".join(paragraphs)
    except Exception as e:
        logger.error(f"DOCX extraction failed for {file_path}: {e}")
        return None


def extract_text(file_path: str) -> dict:
    """
    Extract text from a resume file (PDF or DOCX).
    Returns dict with 'text', 'success', 'error', 'file_type'.
    """
    if not os.path.exists(file_path):
        return {"success": False, "error": "File not found", "text": None, "file_type": None}

    ext = os.path.splitext(file_path)[1].lower()

    if ext == ".pdf":
        text = extract_text_from_pdf(file_path)
        file_type = "pdf"
    elif ext == ".docx":
        text = extract_text_from_docx(file_path)
        file_type = "docx"
    else:
        return {
            "success": False,
            "error": f"Unsupported file type: {ext}",
            "text": None,
            "file_type": None,
        }

    if not text or len(text.strip()) < 50:
        return {
            "success": False,
            "error": "Could not extract sufficient text from the file. The file may be image-based or corrupted.",
            "text": text,
            "file_type": file_type,
        }

    cleaned = _clean_text(text)
    return {"success": True, "error": None, "text": cleaned, "file_type": file_type}


def _clean_text(text: str) -> str:
    """Clean extracted text — normalize whitespace, remove garbage chars."""
    # Normalize line endings
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    # Remove excessive blank lines (keep max 2 consecutive)
    text = re.sub(r"\n{3,}", "\n\n", text)
    # Remove non-printable chars except newline/tab
    text = re.sub(r"[^\x09\x0A\x20-\x7E\u00A0-\uFFFF]", " ", text)
    # Collapse multiple spaces
    text = re.sub(r"[ \t]{2,}", " ", text)
    return text.strip()
